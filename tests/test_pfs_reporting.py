import unittest
import tempfile
from pathlib import Path

from pfs_agent.reporting import (
    AnalysisRequest,
    MetricContract,
    ReportingContractError,
    analyze_csv,
    analyze_file,
    list_xlsx_worksheets,
)


FIXTURE = Path(__file__).resolve().parents[1] / "data" / "fixtures" / "pfs_sales.csv"
METRIC = MetricContract(
    metric_id="sales_amount",
    label="销售额",
    formula="SUM(sales_amount)",
    value_column="sales_amount",
    date_column="month",
    dimension="region",
)


class PfsReportingTests(unittest.TestCase):
    def test_empty_csv_and_header_only_csv_have_distinct_contract_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            empty = Path(directory) / "empty.csv"
            header_only = Path(directory) / "header-only.csv"
            empty.write_text("", encoding="utf-8")
            header_only.write_text("month,region,sales_amount\n", encoding="utf-8")

            with self.assertRaisesRegex(ReportingContractError, "non-empty header") as empty_error:
                analyze_file(empty, metric=METRIC, request=self._request("empty"))
            with self.assertRaisesRegex(ReportingContractError, "at least one data row") as row_error:
                analyze_file(header_only, metric=METRIC, request=self._request("header-only"))

        self.assertEqual("source_header_missing", empty_error.exception.code)
        self.assertEqual("source_has_no_rows", row_error.exception.code)

    def _request(self, run_id, **overrides):
        values = {
            "run_id": run_id,
            "metric_id": METRIC.metric_id,
            "dimension": METRIC.dimension,
            **overrides,
        }
        return AnalysisRequest(**values)

    def test_fixture_produces_stable_grouped_result_and_evidence(self):
        result = analyze_csv(
            FIXTURE,
            metric=METRIC,
            request=AnalysisRequest(
                run_id="run-fixture-001",
                metric_id="sales_amount",
                dimension="region",
                date_from="2026-01",
                date_to="2026-03",
            ),
            source_id="pfs-fixture-sales",
        )
        self.assertEqual("completed", result.status)
        self.assertEqual(100000, result.total)
        self.assertEqual(
            ["华东", "华南", "华北"],
            [item["dimension"] for item in result.groups],
        )
        self.assertEqual([42000, 33000, 25000], [item["value"] for item in result.groups])
        self.assertEqual(2, len(result.claims))
        self.assertEqual(1, len(result.evidence))
        self.assertTrue(result.claims[0]["evidence_ids"])
        self.assertEqual(
            result.evidence[0].content_sha256,
            result.snapshot.content_sha256,
        )

    def test_date_filter_changes_total_without_changing_contract(self):
        result = analyze_csv(
            FIXTURE,
            metric=METRIC,
            request=AnalysisRequest(
                run_id="run-fixture-002",
                metric_id="sales_amount",
                dimension="region",
                date_from="2026-02",
                date_to="2026-02",
            ),
            source_id="pfs-fixture-sales",
        )
        self.assertEqual(33000, result.total)
        self.assertEqual("华东", result.groups[0]["dimension"])
        self.assertEqual(14000, result.groups[0]["value"])

    @unittest.skipUnless(__import__("importlib.util").util.find_spec("openpyxl"), "openpyxl is not installed")
    def test_xlsx_uses_the_same_metric_and_evidence_contract(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sales.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["month", "region", "sales_amount"])
            sheet.append(["2026-01", "华东", 42000])
            sheet.append(["2026-02", "华南", 33000])
            workbook.save(path)
            result = analyze_file(
                path,
                metric=METRIC,
                request=AnalysisRequest(
                    run_id="run-xlsx-001",
                    metric_id="sales_amount",
                    dimension="region",
                ),
                source_id="xlsx-sales",
            )

        self.assertEqual(75000, result.total)
        self.assertEqual(2, result.snapshot.row_count)
        self.assertEqual("xlsx-sales", result.snapshot.source_id)
        self.assertEqual("xlsx", result.snapshot.file_name.rsplit(".", 1)[-1])
        self.assertEqual(result.evidence[0].content_sha256, result.snapshot.content_sha256)

    @unittest.skipUnless(__import__("importlib.util").util.find_spec("openpyxl"), "openpyxl is not installed")
    def test_multi_sheet_xlsx_requires_and_records_explicit_worksheet(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "multi-sheet.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "说明"
            first.append(["note"])
            first.append(["这不是分析数据"])
            sales = workbook.create_sheet("销售明细")
            sales.append(["month", "region", "sales_amount"])
            sales.append(["2026-01", "华东", 1200])
            workbook.save(path)

            self.assertEqual(["说明", "销售明细"], list_xlsx_worksheets(path))
            with self.assertRaisesRegex(ReportingContractError, "choose a worksheet") as error:
                analyze_file(path, metric=METRIC, request=self._request("sheet-required"))
            result = analyze_file(
                path,
                metric=METRIC,
                request=self._request("sheet-selected"),
                worksheet="销售明细",
            )

        self.assertEqual("worksheet_required", error.exception.code)
        self.assertEqual("销售明细", result.snapshot.worksheet)
        self.assertEqual(1200, result.total)

    def test_non_numeric_value_and_no_matching_date_are_explicit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sales.csv"
            path.write_text(
                "month,region,sales_amount\n2026-01,华东,not-a-number\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ReportingContractError, "not numeric") as error:
                analyze_file(path, metric=METRIC, request=self._request("invalid-number"))

            path.write_text(
                "month,region,sales_amount\n2026-01,华东,1200\n",
                encoding="utf-8",
            )
            result = analyze_file(
                path,
                metric=METRIC,
                request=self._request("no-match", date_from="2027-01", date_to="2027-12"),
            )

        self.assertEqual("metric_value_not_numeric", error.exception.code)
        self.assertEqual("unverified", result.status)
        self.assertIn("筛选条件下没有匹配的数据行。", result.warnings)


if __name__ == "__main__":
    unittest.main()
