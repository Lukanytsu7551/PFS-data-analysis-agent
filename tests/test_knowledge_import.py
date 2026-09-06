from io import BytesIO
import os
import tempfile
from types import SimpleNamespace
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from docx import Document

from Function.Knowledge.file_parser import parse_file


class _FakeCompletions:
    def __init__(self, payload: str):
        self.payload = payload
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        # Stream in multiple fragments so the test covers the same assembly
        # path used by an OpenAI-compatible provider.
        pieces = [self.payload[:17], self.payload[17:53], self.payload[53:]]
        return [
            SimpleNamespace(
                choices=[
                    SimpleNamespace(delta=SimpleNamespace(content=piece))
                ]
            )
            for piece in pieces
            if piece
        ]


class _FakeLLMClient:
    def __init__(self, payload: str):
        self.chat = SimpleNamespace(
            completions=_FakeCompletions(payload),
        )
        self.options = {}

    def with_options(self, **kwargs):
        self.options = kwargs
        return self


def _structured_workbook_bytes() -> bytes:
    with tempfile.TemporaryDirectory(prefix="pfs-knowledge-workbook-") as raw:
        path = Path(raw) / "knowledge.xlsx"
        writer = pd.ExcelWriter(path, engine="openpyxl")
        pd.DataFrame([
            {
                "指标名称": "净收入",
                "别名": "Net Revenue",
                "定义": "扣除退款后的收入",
                "SQL模板": "SELECT SUM(net_revenue) FROM orders",
                "备注": "按订单日汇总",
            }
        ]).to_excel(writer, index=False, sheet_name="指标定义")
        pd.DataFrame([
            {
                "规则ID": "revenue_nonnegative",
                "描述": "净收入不可为负",
                "违反条件": "net_revenue < 0",
                "严重程度": "error",
            }
        ]).to_excel(writer, index=False, sheet_name="业务规则")
        pd.DataFrame([
            {
                "主题": "退款分析",
                "内容": "退款率上升时同时检查履约和支付失败",
                "标签": "退款,履约",
            }
        ]).to_excel(writer, index=False, sheet_name="背景知识")
        writer.close()
        return path.read_bytes()


def _mixed_workbook_bytes() -> bytes:
    with tempfile.TemporaryDirectory(prefix="pfs-knowledge-mixed-") as raw:
        path = Path(raw) / "knowledge.xlsx"
        writer = pd.ExcelWriter(path, engine="openpyxl")
        pd.DataFrame([
            {
                "指标名称": "净收入",
                "定义": "扣除退款后的收入",
            }
        ]).to_excel(writer, index=False, sheet_name="指标定义")
        pd.DataFrame([
            {
                "栏目": "经营说明",
                "补充内容": "退款率连续两周上升时，需要同时检查履约和支付失败。",
            }
        ]).to_excel(writer, index=False, sheet_name="说明")
        writer.close()
        return path.read_bytes()


def _docx_bytes() -> bytes:
    with tempfile.TemporaryDirectory(prefix="pfs-knowledge-docx-") as raw:
        path = Path(raw) / "knowledge.docx"
        document = Document()
        document.add_paragraph("退款率上升时同时检查履约和支付失败。")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "指标"
        table.rows[0].cells[1].text = "净收入 = 收入 - 退款"
        document.save(path)
        return path.read_bytes()


def _llm_payload() -> str:
    return (
        '{"metrics":[{"name":"净收入","alias":"Net Revenue",'
        '"definition":"收入减去退款","sql_template":"",'
        '"notes":""}],"business_rules":[],"context_notes":['
        '{"topic":"退款分析","content":"退款率上升时同时检查履约和支付失败",'
        '"tags":"退款,履约"}]}'
    )


class KnowledgeImportTests(unittest.TestCase):
    def test_structured_workbook_parses_without_an_llm(self):
        with tempfile.TemporaryDirectory(prefix="pfs-knowledge-parse-") as raw:
            path = Path(raw) / "knowledge.xlsx"
            path.write_bytes(_structured_workbook_bytes())

            result = parse_file(path, client=None, model="")

        self.assertEqual("structured", result["format"])
        self.assertEqual(3, len(result["preview"]))
        self.assertEqual(
            {"metrics", "business_rules", "context_notes"},
            {item["table"] for item in result["preview"]},
        )

    def test_api_does_not_request_model_for_structured_workbook(self):
        from api import create_app
        from api import knowledge as knowledge_api

        with tempfile.TemporaryDirectory(prefix="pfs-knowledge-api-") as raw:
            with patch.dict(
                os.environ,
                {"PFS_DATA_DIR": str(Path(raw) / "data"), "PFS_NO_BROWSER": "1"},
                clear=False,
            ), patch.object(
                knowledge_api,
                "_get_client",
                side_effect=AssertionError("structured import must not request an LLM"),
            ):
                app = create_app()
                app.config.update(TESTING=True)
                with app.test_client() as client:
                    response = client.post(
                        "/api/knowledge/parse",
                        data={
                            "file": (
                                BytesIO(_structured_workbook_bytes()),
                                "knowledge.xlsx",
                            )
                        },
                        content_type="multipart/form-data",
                    )

        self.assertEqual(200, response.status_code)
        payload = response.get_json()
        self.assertEqual("structured", payload["format"])
        self.assertEqual(3, len(payload["preview"]))

    def test_docx_uses_streaming_llm_extraction_for_paragraphs_and_tables(self):
        client = _FakeLLMClient(_llm_payload())
        with tempfile.TemporaryDirectory(prefix="pfs-knowledge-docx-parse-") as raw:
            path = Path(raw) / "knowledge.docx"
            path.write_bytes(_docx_bytes())

            result = parse_file(path, client=client, model="local-fixture")

        self.assertEqual("unstructured", result["format"])
        self.assertEqual(2, len(result["preview"]))
        self.assertEqual(
            {"metrics", "context_notes"},
            {item["table"] for item in result["preview"]},
        )
        self.assertEqual("local-fixture", client.chat.completions.calls[0]["model"])
        self.assertEqual(120, client.options["timeout"])
        prompt = client.chat.completions.calls[0]["messages"][0]["content"]
        self.assertIn("退款率上升时同时检查履约和支付失败", prompt)
        self.assertIn("净收入 = 收入 - 退款", prompt)

    def test_mixed_workbook_keeps_structured_records_and_adds_llm_records(self):
        client = _FakeLLMClient(_llm_payload())
        with tempfile.TemporaryDirectory(prefix="pfs-knowledge-mixed-parse-") as raw:
            path = Path(raw) / "knowledge.xlsx"
            path.write_bytes(_mixed_workbook_bytes())

            result = parse_file(path, client=client, model="local-fixture")

        self.assertEqual("mixed", result["format"])
        self.assertEqual(3, len(result["preview"]))
        self.assertEqual(
            {"metrics", "context_notes"},
            {item["table"] for item in result["preview"]},
        )
        self.assertEqual(1, len(client.chat.completions.calls))

    def test_api_requests_model_for_unstructured_docx_only_after_parser_fallback(self):
        from api import create_app
        from api import knowledge as knowledge_api

        fake_client = _FakeLLMClient(_llm_payload())
        with tempfile.TemporaryDirectory(prefix="pfs-knowledge-docx-api-") as raw:
            with patch.dict(
                os.environ,
                {"PFS_DATA_DIR": str(Path(raw) / "data"), "PFS_NO_BROWSER": "1"},
                clear=False,
            ), patch.object(
                knowledge_api,
                "_get_client",
                return_value=(fake_client, "local-fixture"),
            ) as get_client:
                app = create_app()
                app.config.update(TESTING=True)
                with app.test_client() as client:
                    response = client.post(
                        "/api/knowledge/parse",
                        data={
                            "session_id": "knowledge-docx-api",
                            "file": (BytesIO(_docx_bytes()), "knowledge.docx"),
                        },
                        content_type="multipart/form-data",
                    )

        self.assertEqual(200, response.status_code)
        payload = response.get_json()
        self.assertEqual("unstructured", payload["format"])
        self.assertEqual(2, len(payload["preview"]))
        get_client.assert_called_once_with("knowledge-docx-api", provider="")


if __name__ == "__main__":
    unittest.main()
