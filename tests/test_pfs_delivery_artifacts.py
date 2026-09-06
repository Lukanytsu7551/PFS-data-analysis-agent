import io
import json
import os
import subprocess
import sys
import unittest
import uuid
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd

from api import create_app
from api.state import session_manager
from data.workspace import WorkspaceManager
from data.workspace_metadata import WorkspaceMetadataStore
from data.workflow_run_store import WorkflowRunStore, WorkflowRunStoreError
from infrastructure.artifact_lifecycle import (
    list_registered_artifacts,
    prune_registry_for_paths,
    register_artifact,
    update_artifact_run_usage,
)
from infrastructure.paths import data_path
from agent.agent import BusinessAgent
from agent.workflows.runtime import WorkflowRuntime


class PfsDeliveryArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.app.config.update(TESTING=True)

    def setUp(self):
        self.client = self.app.test_client()
        self.sid = f"pfs-delivery-{uuid.uuid4().hex[:12]}"
        session_manager.get_or_create(self.sid)

    def tearDown(self):
        session_manager.remove(self.sid)

    def test_session_artifact_history_is_scoped_and_returns_safe_result_metadata(self):
        other_sid = f"pfs-other-{uuid.uuid4().hex[:12]}"
        session_manager.get_or_create(other_sid)
        try:
            with TemporaryDirectory() as tmp, patch.dict(os.environ, {"PFS_DATA_DIR": tmp}, clear=False):
                output_dir = data_path("outputs", "exports")
                output_dir.mkdir(parents=True, exist_ok=True)
                artifact_path = output_dir / "scoped.xlsx"
                artifact_path.write_bytes(b"pfs artifact")
                artifact_id = register_artifact(
                    artifact_path, artifact_type="xlsx", session_id=self.sid,
                    metadata={"run_id": "scoped-run", "claim_ids": [], "evidence_ids": [],
                              "analysis_parameters": {"metric_id": "sales_amount"}},
                )
                listed = self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts")
                self.assertEqual(200, listed.status_code)
                self.assertEqual(artifact_id, listed.get_json()["artifacts"][0]["id"])
                self.assertEqual("scoped-run", listed.get_json()["artifacts"][0]["run_id"])
                self.assertIn("/download", listed.get_json()["artifacts"][0]["download_url"])
                self.assertNotIn("lineage", listed.get_json()["artifacts"][0])
                detail = self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}")
                self.assertEqual(200, detail.status_code)
                self.assertEqual(artifact_id, detail.get_json()["artifact"]["id"])
                detail_artifact = detail.get_json()["artifact"]
                self.assertNotIn("lineage", detail_artifact)
                self.assertNotIn("claim_ids", detail_artifact)
                self.assertNotIn("evidence_ids", detail_artifact)
                downloaded = self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}/download")
                self.assertEqual(200, downloaded.status_code)
                self.assertEqual(b"pfs artifact", downloaded.data)
                downloaded.close()
                self.assertEqual(1, self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}").get_json()["artifact"]["download_count"])
                cross_session = self.client.get(f"/api/session/{other_sid}/lifecycle/artifacts/{artifact_id}")
                self.assertEqual(404, cross_session.status_code)
                self.assertEqual("artifact_not_found", cross_session.get_json()["code"])
                self.assertEqual(404, self.client.get(f"/api/session/{other_sid}/lifecycle/artifacts/{artifact_id}/download").status_code)
        finally:
            session_manager.remove(other_sid)

    def test_run_usage_updates_only_matching_session_and_run_without_changing_identity(self):
        other_sid = f"pfs-other-{uuid.uuid4().hex[:12]}"
        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ):
            output_dir = data_path("outputs", "exports")
            output_dir.mkdir(parents=True, exist_ok=True)
            paths = [output_dir / name for name in ("run-a.xlsx", "run-a.docx", "run-b.xlsx")]
            for index, path in enumerate(paths):
                path.write_bytes(f"artifact-{index}".encode())
            run_a_ids = [
                register_artifact(
                    path, artifact_type=path.suffix.lstrip("."),
                    session_id=self.sid, metadata={"run_id": "run-a"},
                )
                for path in paths[:2]
            ]
            run_b_id = register_artifact(
                paths[2], artifact_type="xlsx", session_id=self.sid,
                metadata={"run_id": "run-b"},
            )
            before = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts/{run_a_ids[0]}"
            ).get_json()["artifact"]

            cost = {
                "amount": 0.0012, "currency": "USD", "estimated": True,
                "source": "provider_usage_configured_pricing",
                "model_calls": 2, "input_tokens": 100, "output_tokens": 20,
                "cached_input_tokens": 8, "providers": ["deepseek"],
                "models": ["deepseek-chat"], "status": "succeeded",
            }
            self.assertEqual(0, update_artifact_run_usage(
                session_id=other_sid, run_id="run-a", cost=cost,
            ))
            self.assertEqual(2, update_artifact_run_usage(
                session_id=self.sid, run_id="run-a", cost=cost,
            ))

            listed = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts"
            ).get_json()["artifacts"]
            indexed = {item["id"]: item for item in listed}
            for artifact_id in run_a_ids:
                self.assertEqual(2, indexed[artifact_id]["cost"]["model_calls"])
                self.assertEqual(0.0012, indexed[artifact_id]["cost"]["amount"])
                self.assertFalse(indexed[artifact_id]["cost"]["billing_verified"])
            self.assertNotIn("cost", indexed[run_b_id])

            after = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts/{run_a_ids[0]}"
            ).get_json()["artifact"]
            for field in ("id", "sha256", "size_bytes", "run_id", "workspace_id"):
                self.assertEqual(before.get(field), after.get(field))

    def test_workflow_export_links_run_cost_and_refreshes_terminal_total(self):
        class FakeRunStore:
            def __init__(self):
                self.nodes = [{
                    "model_name": "deepseek-chat", "provider_name": "deepseek",
                    "model_calls": 1, "input_tokens": 80, "output_tokens": 20,
                    "cached_input_tokens": 8, "cost_usd": 0.001,
                }]
                self.events = []

            def list_node_runs(self, _run_id):
                return list(self.nodes)

            def record_event(self, run_id, event_type, payload):
                self.events.append((run_id, event_type, payload))

        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ):
            artifact_dir = data_path("outputs", "workflow")
            artifact_dir.mkdir(parents=True, exist_ok=True)
            runtime = WorkflowRuntime.__new__(WorkflowRuntime)
            runtime.session_id = self.sid
            runtime.workspace = type("Workspace", (), {
                "artifacts_dir": artifact_dir,
                "workspace_id": "workflow-cost-workspace",
            })()
            runtime.run_store = FakeRunStore()
            result = runtime._execute_export_node(
                {
                    "node_id": "export", "output_contract": ["delivery"],
                    "export": {"source": "report", "format": "markdown"},
                    "__pfs_workflow_context__": {
                        "run_id": "workflow-run-1", "node_run_id": "node-run-export",
                    },
                },
                {"report": "# report"},
            )["delivery"]

            artifacts = list_registered_artifacts(session_id=self.sid)
            self.assertEqual(result["artifact_id"], artifacts[0]["id"])
            self.assertEqual("workflow-run-1", artifacts[0]["run_id"])
            self.assertEqual("running", artifacts[0]["cost"]["status"])
            self.assertEqual(0.001, artifacts[0]["cost"]["amount"])
            self.assertEqual(1, artifacts[0]["cost"]["model_calls"])

            runtime.run_store.nodes.append({
                "model_name": "deepseek-chat", "provider_name": "deepseek",
                "model_calls": 2, "input_tokens": 40, "output_tokens": 10,
                "cached_input_tokens": 0, "cost_usd": 0.0005,
            })
            runtime._finalize_workflow_artifact_cost("workflow-run-1", "succeeded")
            finalized = list_registered_artifacts(session_id=self.sid)[0]
            self.assertEqual("succeeded", finalized["cost"]["status"])
            self.assertEqual(0.0015, finalized["cost"]["amount"])
            self.assertEqual(3, finalized["cost"]["model_calls"])

            runtime.run_store.nodes[1]["cost_usd"] = None
            unknown = runtime._workflow_run_cost(
                "workflow-run-1", status="succeeded",
            )
            self.assertIsNone(unknown["amount"])
            self.assertEqual("provider_usage_price_unknown", unknown["source"])

    def test_workflow_export_reuses_completed_side_effect_without_duplicate_artifact(self):
        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ):
            workspace_id = "workflow-idempotency-workspace"
            db_path = Path(tmp) / "workflow.sqlite3"
            run_store = WorkflowRunStore(db_path, workspace_id)
            try:
                graph = {
                    "nodes": [{"node_id": "export", "type": "export"}],
                    "edges": [],
                    "entry_node_ids": ["export"],
                }
                run = run_store.create_run(
                    workflow_version_id="workflow-version",
                    session_id=self.sid,
                    graph=graph,
                    inputs={},
                )
                node_run = run_store.list_node_runs(run["id"])[0]
                artifact_dir = data_path("outputs", "workflow")
                runtime = WorkflowRuntime.__new__(WorkflowRuntime)
                runtime.session_id = self.sid
                runtime.workspace = type("Workspace", (), {
                    "artifacts_dir": artifact_dir,
                    "workspace_id": workspace_id,
                })()
                runtime.run_store = run_store
                node = {
                    "node_id": "export",
                    "output_contract": ["delivery"],
                    "export": {"source": "report", "format": "markdown"},
                    "__pfs_workflow_context__": {
                        "run_id": run["id"],
                        "node_run_id": node_run["id"],
                        "node_id": "export",
                    },
                }
                first = runtime._execute_export_node(node, {"report": "# idempotent report"})
                second = runtime._execute_export_node(node, {"report": "# idempotent report"})

                self.assertEqual(first, second)
                self.assertEqual(1, len(list_registered_artifacts(session_id=self.sid)))
                with self.assertRaises(WorkflowRunStoreError):
                    runtime._execute_export_node(node, {"report": "# changed payload"})
                effect = run_store.get_side_effect_for_node_run(
                    node_run["id"], effect_type="export_file",
                )
                self.assertEqual("succeeded", effect["status"])
                self.assertEqual(first, effect["result"])
                event_types = [item["type"] for item in run_store.list_events(run["id"])]
                self.assertIn("workflow_side_effect_claimed", event_types)
                self.assertIn("workflow_side_effect_completed", event_types)
            finally:
                run_store.close()

    def test_unknown_provider_price_keeps_tokens_but_never_fakes_zero_cost(self):
        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ):
            output_dir = data_path("outputs", "exports")
            output_dir.mkdir(parents=True, exist_ok=True)
            artifact_path = output_dir / "unknown-price.xlsx"
            artifact_path.write_bytes(b"unknown-price")
            artifact_id = register_artifact(
                artifact_path, artifact_type="xlsx", session_id=self.sid,
                metadata={"run_id": "unknown-price-run"},
            )
            self.assertEqual(1, update_artifact_run_usage(
                session_id=self.sid, run_id="unknown-price-run",
                cost={
                    "amount": None, "currency": "USD", "estimated": None,
                    "source": "provider_usage_price_unknown",
                    "model_calls": 1, "input_tokens": 80, "output_tokens": 20,
                    "cached_input_tokens": 0, "providers": ["deepseek"],
                    "models": ["deepseek-chat"], "status": "succeeded",
                },
            ))
            artifact = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}"
            ).get_json()["artifact"]
            self.assertIsNone(artifact["cost"]["amount"])
            self.assertIsNone(artifact["cost"]["estimated"])
            self.assertEqual(100, artifact["cost"]["input_tokens"] + artifact["cost"]["output_tokens"])
            with self.assertRaisesRegex(ValueError, "必须保持 null"):
                update_artifact_run_usage(
                    session_id=self.sid, run_id="unknown-price-run",
                    cost={**artifact["cost"], "amount": 0.0},
                )

    def test_chat_run_usage_is_persisted_per_turn_and_shared_by_turn_artifacts(self):
        captured_runs = []

        class FakeAgent:
            def __init__(self, session_id):
                self._provider = "deepseek"
                self.model = "deepseek-chat"
                self._artifact_metadata = {}
                self._session_id = session_id

            def run(self, _message, _history, **kwargs):
                run_id = kwargs["run_id"]
                captured_runs.append(run_id)
                for suffix in ("xlsx", "docx"):
                    path = data_path("outputs", "exports", f"{run_id}.{suffix}")
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_bytes(f"{run_id}-{suffix}".encode())
                    register_artifact(
                        path, artifact_type=suffix, session_id=self._session_id,
                        metadata={**self._artifact_metadata, "run_id": run_id},
                    )
                yield {
                    "type": "usage", "run_id": run_id,
                    "provider": self._provider, "model": self.model,
                    "model_calls": 1, "prompt_tokens": 60,
                    "completion_tokens": 10, "cached_input_tokens": 4,
                    "cost_usd": 0.0007,
                }
                yield {
                    "type": "usage", "run_id": run_id,
                    "provider": self._provider, "model": self.model,
                    "model_calls": 1, "prompt_tokens": 40,
                    "completion_tokens": 10, "cached_input_tokens": 0,
                    "cost_usd": 0.0005,
                }
                yield {"type": "text", "content": "分析完成"}
                yield {"type": "done"}

        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ), patch(
            "api.chat._build_agent",
            side_effect=lambda *_args, **_kwargs: FakeAgent(self.sid),
        ):
            for message in ("第一轮", "第二轮"):
                response = self.client.post(
                    f"/api/session/{self.sid}/chat", json={"message": message},
                )
                self.assertEqual(200, response.status_code)
                self.assertIn("分析完成", response.get_data(as_text=True))

            self.assertEqual(2, len(captured_runs))
            self.assertNotEqual(captured_runs[0], captured_runs[1])
            artifacts = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts?limit=20"
            ).get_json()["artifacts"]
            by_run = {}
            for artifact in artifacts:
                by_run.setdefault(artifact.get("run_id"), []).append(artifact)
            for run_id in captured_runs:
                self.assertEqual(2, len(by_run[run_id]))
                for artifact in by_run[run_id]:
                    self.assertEqual(2, artifact["cost"]["model_calls"])
                    self.assertEqual(100, artifact["cost"]["input_tokens"])
                    self.assertEqual(20, artifact["cost"]["output_tokens"])
                    self.assertEqual(0.0012, artifact["cost"]["amount"])
                    self.assertEqual("succeeded", artifact["cost"]["status"])

    def test_registered_artifact_recycle_restore_is_recoverable_and_idempotent(self):
        with TemporaryDirectory() as tmp, patch.dict(os.environ, {"PFS_DATA_DIR": tmp}, clear=False):
            output_dir = data_path("outputs", "exports")
            output_dir.mkdir(parents=True, exist_ok=True)
            artifact_path = output_dir / "restore-me.xlsx"
            artifact_path.write_bytes(b"recoverable pfs artifact")
            artifact_id = register_artifact(
                artifact_path, artifact_type="xlsx", session_id=self.sid,
                metadata={
                    "run_id": "restore-run",
                    "source_sha256": "source-snapshot",
                    "claim_ids": ["claim-1"],
                    "evidence_ids": ["evidence-1"],
                    "cost": {
                        "amount": 0.0, "currency": "USD", "estimated": False,
                        "source": "deterministic_no_model", "model_calls": 0,
                    },
                },
            )
            downloaded = self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}/download")
            self.assertEqual(200, downloaded.status_code)
            downloaded.close()
            recycled = self.client.post(
                "/api/lifecycle/artifacts/registered/recycle",
                json={"artifact_id": artifact_id},
            )
            self.assertEqual(200, recycled.status_code, recycled.get_data(as_text=True))
            trash_id = recycled.get_json()["summary"]["trash_id"]
            self.assertEqual([], self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts").get_json()["artifacts"])
            restored = self.client.post(f"/api/lifecycle/artifact-trash/{trash_id}/restore")
            self.assertEqual(200, restored.status_code, restored.get_data(as_text=True))
            self.assertEqual(artifact_id, restored.get_json()["summary"]["artifact_id"])
            restored_items = self.client.get(f"/api/session/{self.sid}/lifecycle/artifacts").get_json()["artifacts"]
            self.assertEqual([artifact_id], [item["id"] for item in restored_items])
            restored_item = restored_items[0]
            self.assertEqual("restore-run", restored_item["run_id"])
            self.assertEqual("source-snapshot", restored_item["source_sha256"])
            self.assertEqual(1, restored_item["download_count"])
            self.assertEqual("deterministic_no_model", restored_item["cost"]["source"])
            self.assertTrue(artifact_path.is_file())

            probe = subprocess.run(
                [
                    sys.executable,
                    "-c",
                    (
                        "import json; "
                        "from api import create_app; "
                        "app=create_app(); app.config.update(TESTING=True); client=app.test_client(); "
                        f"detail=client.get('/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}'); "
                        f"download=client.get('/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}/download'); "
                        "print(json.dumps({'status': detail.status_code, 'artifact': detail.get_json()['artifact'], "
                        "'download_status': download.status_code, 'download_body': download.data.decode()}))"
                    ),
                ],
                cwd=Path(__file__).resolve().parents[1],
                env={**os.environ, "PFS_DATA_DIR": tmp},
                capture_output=True,
                text=True,
                check=True,
            )
            process_readback = json.loads(probe.stdout)
            self.assertEqual(200, process_readback["status"])
            self.assertEqual(200, process_readback["download_status"])
            self.assertEqual("recoverable pfs artifact", process_readback["download_body"])
            process_item = process_readback["artifact"]
            self.assertEqual(artifact_id, process_item["id"])
            self.assertEqual("restore-run", process_item["run_id"])
            self.assertEqual(1, process_item["download_count"])
            self.assertNotIn("claim_ids", process_item)
            self.assertNotIn("evidence_ids", process_item)
            repeated = self.client.post(f"/api/lifecycle/artifact-trash/{trash_id}/restore")
            self.assertEqual(404, repeated.status_code)

    def test_workspace_artifact_keeps_stable_identity_after_switch_and_process_reopen(self):
        with TemporaryDirectory() as data_tmp, TemporaryDirectory() as workspace_tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": data_tmp}, clear=False,
        ):
            workspace_root = Path(workspace_tmp).resolve()
            store = WorkspaceMetadataStore()
            manager = WorkspaceManager(metadata_store=store, remember_mounts_by_default=True)
            ok, message, runtime = manager.mount(
                self.sid, str(workspace_root), permission="read_write",
            )
            self.assertTrue(ok, message)
            artifact_path = runtime.artifacts_dir / "workspace-report.xlsx"
            artifact_path.write_bytes(b"workspace-linked-artifact")

            with patch("data.workspace.workspace_manager", manager):
                artifact_id = register_artifact(
                    artifact_path,
                    artifact_type="xlsx",
                    session_id=self.sid,
                    workspace_id=runtime.workspace_id,
                    artifact_id="workspace-artifact-stable",
                    metadata={"run_id": "workspace-run"},
                )
                initial = self.client.get(
                    f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}"
                ).get_json()["artifact"]
            self.assertEqual(runtime.workspace_id, initial["workspace_id"])
            self.assertEqual(workspace_root.name, initial["workspace"]["name"])
            self.assertNotIn(str(workspace_root), json.dumps(initial, ensure_ascii=False))

            self.assertTrue(manager.unmount(self.sid))
            reopened_manager = WorkspaceManager(
                metadata_store=WorkspaceMetadataStore(),
                remember_mounts_by_default=True,
            )
            with patch("data.workspace.workspace_manager", reopened_manager):
                reopened = self.client.get(
                    f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}"
                )
                downloaded = self.client.get(
                    f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}/download"
                )
            self.assertEqual(200, reopened.status_code)
            reopened_artifact = reopened.get_json()["artifact"]
            self.assertEqual(runtime.workspace_id, reopened_artifact["workspace_id"])
            self.assertEqual(workspace_root.name, reopened_artifact["workspace"]["name"])
            self.assertTrue(reopened_artifact["workspace"]["available"])
            self.assertEqual(200, downloaded.status_code)
            self.assertEqual(b"workspace-linked-artifact", downloaded.data)
            downloaded.close()

            probe = subprocess.run(
                [
                    sys.executable,
                    "-c",
                    (
                        "import json; from api import create_app; "
                        "app=create_app(); app.config.update(TESTING=True); client=app.test_client(); "
                        f"response=client.get('/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}'); "
                        "print(json.dumps({'status': response.status_code, 'artifact': response.get_json()['artifact']}))"
                    ),
                ],
                cwd=Path(__file__).resolve().parents[1],
                env={**os.environ, "PFS_DATA_DIR": data_tmp},
                capture_output=True,
                text=True,
                check=True,
            )
            process_item = json.loads(probe.stdout)
            self.assertEqual(200, process_item["status"])
            self.assertEqual(runtime.workspace_id, process_item["artifact"]["workspace_id"])
            self.assertEqual(workspace_root.name, process_item["artifact"]["workspace"]["name"])

            with patch("data.workspace.workspace_manager", reopened_manager):
                duplicate_id = register_artifact(
                    artifact_path,
                    artifact_type="xlsx",
                    session_id=self.sid,
                    workspace_id=runtime.workspace_id,
                    artifact_id=artifact_id,
                    metadata={"run_id": "workspace-run"},
                )
            self.assertEqual(artifact_id, duplicate_id)
            with patch("data.workspace.workspace_manager", reopened_manager):
                listed = self.client.get(
                    f"/api/session/{self.sid}/lifecycle/artifacts"
                ).get_json()["artifacts"]
            self.assertEqual(1, len([item for item in listed if item["id"] == artifact_id]))

            with patch("data.workspace.workspace_manager", reopened_manager):
                recycled = self.client.post(
                    "/api/lifecycle/artifacts/registered/recycle",
                    json={"artifact_id": artifact_id},
                )
                self.assertEqual(200, recycled.status_code, recycled.get_data(as_text=True))
                self.assertFalse(artifact_path.exists())
                trash_id = recycled.get_json()["summary"]["trash_id"]
                restored = self.client.post(
                    f"/api/lifecycle/artifact-trash/{trash_id}/restore"
                )
            self.assertEqual(200, restored.status_code, restored.get_data(as_text=True))
            self.assertEqual(artifact_id, restored.get_json()["summary"]["artifact_id"])
            self.assertEqual(b"workspace-linked-artifact", artifact_path.read_bytes())

            # The managed data-root sweeper may encounter the same relative
            # filename. It must never prune a workspace-owned registry entry.
            self.assertEqual(0, prune_registry_for_paths({"artifacts/workspace-report.xlsx"}))
            with patch("data.workspace.workspace_manager", reopened_manager):
                retained = self.client.get(
                    f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}"
                )
            self.assertEqual(200, retained.status_code)

    def test_fixture_delivery_generates_parseable_office_files(self):
        with TemporaryDirectory() as tmp, patch.object(
            BusinessAgent, "_get_export_dir", return_value=tmp
        ), patch("agent.tools.business.export.register_artifact"):
            for output_format, extension in (("xlsx", ".xlsx"), ("docx", ".docx"), ("pptx", ".pptx")):
                with self.subTest(output_format=output_format):
                    response = self.client.post(
                        "/api/pfs/deliver",
                        json={"format": output_format, "run_id": f"delivery-{output_format}"},
                    )
                    self.assertEqual(200, response.status_code, response.get_data(as_text=True))
                    payload = response.get_json()
                    self.assertTrue(payload["ok"])
                    artifact = payload["artifacts"][0]
                    self.assertEqual(output_format, artifact["type"])
                    self.assertEqual(f"delivery-{output_format}", artifact["run_id"])
                    self.assertEqual({
                        "amount": 0.0,
                        "currency": "USD",
                        "estimated": False,
                        "source": "deterministic_no_model",
                        "model_calls": 0,
                        "input_tokens": 0,
                        "output_tokens": 0,
                    }, artifact["cost"])
                    self.assertTrue(artifact["source_sha256"])
                    self.assertIn("analysis_parameters", artifact)
                    self.assertTrue(artifact["sql"])
                    self.assertEqual(2, len(artifact["chart_specs"]))
                    self.assertEqual(2, len(artifact["final_claims"]))
                    self.assertIn("metric_id", artifact["analysis_parameters"])
                    self.assertTrue(artifact["name"].endswith(extension))
                    self.assertTrue((Path(tmp) / artifact["name"]).is_file())

            from openpyxl import load_workbook
            from docx import Document
            from pptx import Presentation

            workbook = load_workbook(next(Path(tmp).glob("*.xlsx")), read_only=True, data_only=True)
            self.assertIn("pfs_sales", workbook.sheetnames)
            workbook.close()
            document_text = "\n".join(
                paragraph.text for paragraph in Document(next(Path(tmp).glob("*.docx"))).paragraphs
            )
            self.assertIn("PFS 销售额分析文档", document_text)
            self.assertIn("SHA-256", document_text)
            presentation = Presentation(next(Path(tmp).glob("*.pptx")))
            self.assertEqual(4, len(presentation.slides))

    def test_uploaded_delivery_artifact_detail_keeps_safe_result_metadata(self):
        csv = (
            "month,region,sales_amount\n"
            "2026-04,华东,2100\n"
            "2026-04,华南,900\n"
        ).encode("utf-8")
        with TemporaryDirectory() as tmp, patch.dict(
            os.environ, {"PFS_DATA_DIR": tmp}, clear=False,
        ):
            uploaded = self.client.post(
                f"/api/session/{self.sid}/upload",
                data={"file": (io.BytesIO(csv), "uploaded-lineage.csv")},
                content_type="multipart/form-data",
            )
            self.assertEqual(200, uploaded.status_code, uploaded.get_data(as_text=True))
            source_id = uploaded.get_json()["added"][0]["source_id"]

            response = self.client.post(
                f"/api/session/{self.sid}/pfs/deliver",
                json={
                    "format": "xlsx",
                    "source_id": source_id,
                    "run_id": "uploaded-lineage-run",
                    "value_column": "sales_amount",
                    "date_column": "month",
                    "dimension": "region",
                },
            )
            self.assertEqual(200, response.status_code, response.get_data(as_text=True))
            artifact_id = response.get_json()["artifacts"][0]["artifact_id"]

            detail_response = self.client.get(
                f"/api/session/{self.sid}/lifecycle/artifacts/{artifact_id}"
            )
            self.assertEqual(200, detail_response.status_code, detail_response.get_data(as_text=True))
            artifact = detail_response.get_json()["artifact"]
            self.assertEqual("uploaded-lineage-run", artifact["run_id"])
            self.assertEqual(2, artifact["included_rows"])
            self.assertTrue(artifact["source_sha256"])
            self.assertTrue(artifact.get("final_claims"))
            self.assertNotIn("lineage", artifact)
            self.assertNotIn("claim_ids", artifact)
            self.assertNotIn("evidence_ids", artifact)

    def test_fixture_delivery_generates_dashboard_links_and_pfs_content(self):
        import api.dashboard as dashboard_module

        with TemporaryDirectory() as tmp, patch.object(
            dashboard_module, "_DASHBOARD_DIR", tmp
        ):
            response = self.client.post(
                "/api/pfs/deliver",
                json={"format": "dashboard", "run_id": "delivery-dashboard"},
            )
            self.assertEqual(200, response.status_code, response.get_data(as_text=True))
            payload = response.get_json()
            self.assertTrue(payload["ok"])
            self.assertEqual(["dashboard", "dashboard"], [item["type"] for item in payload["artifacts"]])
            dashboard_files = list(Path(tmp).glob("*.json"))
            self.assertEqual(1, len(dashboard_files))
            dashboard = json.loads(dashboard_files[0].read_text(encoding="utf-8"))
            self.assertEqual("PFS 销售额分析看板", dashboard["name"])
            self.assertEqual(2, len(dashboard["widgets"]))

            export_response = self.client.get(payload["artifacts"][1]["url"])
            self.assertEqual(200, export_response.status_code)
            disposition = export_response.headers["Content-Disposition"]
            disposition.encode("latin-1")
            self.assertIn("filename*=UTF-8''", disposition)
            self.assertIn("PFS_", disposition)
            export_response.close()

    def test_delivery_rejects_unknown_format(self):
        response = self.client.post(
            "/api/pfs/deliver",
            json={"format": "pdf", "run_id": "delivery-pdf"},
        )
        self.assertEqual(400, response.status_code)
        self.assertFalse(response.get_json()["ok"])
        self.assertIn("xlsx, docx, pptx, or dashboard", response.get_json()["error"])

    def test_delivery_wraps_unexpected_export_failure_with_stable_code(self):
        with patch.object(
            BusinessAgent,
            "_tool_export_excel",
            side_effect=OSError("output directory is not writable"),
        ):
            response = self.client.post(
                "/api/pfs/deliver",
                json={"format": "xlsx", "run_id": "delivery-failure"},
            )

        self.assertEqual(400, response.status_code)
        payload = response.get_json()
        self.assertFalse(payload["ok"])
        self.assertEqual("delivery_generation_failed", payload["code"])
        self.assertEqual("交付物生成失败", payload["error"])
        self.assertNotIn("output directory", payload["error"])

    def test_multi_sheet_delivery_uses_only_the_analyzed_worksheet(self):
        import api.dashboard as dashboard_module

        with TemporaryDirectory() as source_tmp, TemporaryDirectory() as export_tmp, TemporaryDirectory() as dashboard_tmp:
            workbook_path = Path(source_tmp) / "regional workbook.xlsx"
            with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
                pd.DataFrame({"note": ["do not export"]}).to_excel(
                    writer, sheet_name="Read Me", index=False
                )
                pd.DataFrame(
                    {
                        "month": ["2026-01", "2026-02"],
                        "region": ["华东", "华南"],
                        "sales_amount": [1200, 900],
                    }
                ).to_excel(writer, sheet_name="Sales Data", index=False)
            with workbook_path.open("rb") as handle:
                uploaded = self.client.post(
                    f"/api/session/{self.sid}/upload",
                    data={"file": (handle, workbook_path.name)},
                    content_type="multipart/form-data",
                )
            source_id = uploaded.get_json()["added"][0]["source_id"]
            report = {
                "source_id": source_id,
                "worksheet": "Sales Data",
                "value_column": "sales_amount",
                "date_column": "month",
                "dimension": "region",
                "run_id": "selected-sheet-delivery",
            }

            with patch.object(BusinessAgent, "_get_export_dir", return_value=export_tmp), patch(
                "agent.tools.business.export.register_artifact"
            ):
                excel_response = self.client.post(
                    f"/api/session/{self.sid}/pfs/deliver", json={**report, "format": "xlsx"}
                )
            self.assertEqual(200, excel_response.status_code, excel_response.get_data(as_text=True))
            from openpyxl import load_workbook

            exported = load_workbook(next(Path(export_tmp).glob("*.xlsx")), read_only=True)
            self.assertEqual(["Sales_Data"], exported.sheetnames)
            exported.close()

            with patch.object(dashboard_module, "_DASHBOARD_DIR", dashboard_tmp):
                dashboard_response = self.client.post(
                    f"/api/session/{self.sid}/pfs/deliver",
                    json={**report, "format": "dashboard"},
                )
            self.assertEqual(
                200, dashboard_response.status_code, dashboard_response.get_data(as_text=True)
            )
            dashboard = json.loads(next(Path(dashboard_tmp).glob("*.json")).read_text(encoding="utf-8"))
            sql = "\n".join(widget["sql"] for widget in dashboard["widgets"])
            self.assertIn('FROM "Sales_Data"', sql)
            self.assertNotIn('FROM "Read_Me"', sql)


if __name__ == "__main__":
    unittest.main()
